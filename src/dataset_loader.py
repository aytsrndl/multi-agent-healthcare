from pathlib import Path

from openpyxl import load_workbook

from src.schemas import Domain, LiteracyLevel, QuestionInput


LITERACY_COLUMN_MAP = {
    LiteracyLevel.VERY_LOW: (
        "Very Low: Inadequate HL + Limited Written English"
    ),
    LiteracyLevel.INADEQUATE: "Inadequate Health Literacy",
    LiteracyLevel.MARGINAL: "Marginal Health Literacy",
    LiteracyLevel.ADEQUATE: "Adequate Health Literacy",
}


DOMAIN_MAP = {
    "Infant Care": Domain.INFANT_CARE,
    "Postpartum Physical Recovery": Domain.PHYSICAL_RECOVERY,
    "Postpartum Mental Health Recovery": Domain.MENTAL_HEALTH,
}


def normalize_category(category: str) -> str:
    """
    Convert spreadsheet category labels such as:

        '1. Infant Care'
        '2. Postpartum Physical Recovery'
        '2. Postpartum Physical Recovery*'

    into clean category names that match DOMAIN_MAP.
    """

    category = str(category).strip()

    # Remove trailing footnote marker
    category = category.replace("*", "")

    # Remove numbering at the beginning
    if ". " in category:
        category = category.split(". ", 1)[1]

    return category.strip()


def load_questions(
    file_path: str | Path,
    literacy_level: LiteracyLevel,
) -> list[QuestionInput]:

    workbook = load_workbook(
        filename=file_path,
        data_only=True,
    )

    if "Results" not in workbook.sheetnames:
        raise ValueError(
            "Expected worksheet 'Results' was not found."
        )

    worksheet = workbook["Results"]

    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(
            worksheet[1],
            start=1,
        )
        if cell.value is not None
    }

    question_column = LITERACY_COLUMN_MAP[literacy_level]

    required_columns = [
        "category",
        "ID",
        "Original Question",
        question_column,
    ]

    for column in required_columns:
        if column not in headers:
            raise ValueError(
                f"Required column '{column}' was not found.\n"
                f"Available columns: {list(headers.keys())}"
            )

    questions = []

    for row in range(2, worksheet.max_row + 1):

        question_id = worksheet.cell(
            row=row,
            column=headers["ID"],
        ).value

        category_text = worksheet.cell(
            row=row,
            column=headers["category"],
        ).value

        question_text = worksheet.cell(
            row=row,
            column=headers[question_column],
        ).value

        if question_id is None:
            continue

        if category_text is None:
            raise ValueError(
                f"Question {question_id} has no category."
            )

        normalized_category = normalize_category(
            category_text
        )

        if normalized_category not in DOMAIN_MAP:
            raise ValueError(
                f"Unknown category for question {question_id}: "
                f"'{category_text}' -> "
                f"'{normalized_category}'"
            )

        if question_text is None:
            raise ValueError(
                f"Question {question_id} has no text for "
                f"{literacy_level.value}."
            )

        question = QuestionInput(
            question_id=int(question_id),
            text=str(question_text).strip(),
            domain=DOMAIN_MAP[normalized_category],
            literacy_level=literacy_level,
        )

        questions.append(question)

    return questions